"""document_loader.py 单元测试：docx 解析、元数据归一与分发逻辑。"""
import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document

import document_loader


def test_load_docx_reads_paragraphs(tmp_path):
    path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    doc.save(str(path))

    docs = document_loader.load_docx(str(path))
    assert len(docs) == 1
    assert "第一段内容" in docs[0].page_content
    assert "第二段内容" in docs[0].page_content
    # 默认 source 为文件名（而非临时完整路径）
    assert docs[0].metadata.get("source") == "sample.docx"


def test_load_docx_skips_empty_paragraphs(tmp_path):
    path = tmp_path / "empty.docx"
    doc = DocxDocument()
    doc.add_paragraph("")  # 空段落应被过滤
    doc.add_paragraph("有内容")
    doc.save(str(path))

    docs = document_loader.load_docx(str(path))
    assert docs[0].page_content == "有内容"


def test_load_docx_uses_source_name_and_no_page(tmp_path):
    path = tmp_path / "x.docx"
    DocxDocument().save(str(path))  # 空文档即可

    docs = document_loader.load_docx(str(path), source_name="说明.docx")
    assert docs[0].metadata.get("source") == "说明.docx"
    # docx 无真实分页，不应设 page
    assert "page" not in docs[0].metadata


def test_load_pdf_tags_source_and_page(tmp_path, monkeypatch):
    """PDF 元数据应归一为友好来源名 + 从 1 起的页码（不依赖真实 PDF）。"""

    class FakeLoader:
        def __init__(self, path):
            self.path = path

        def load(self):
            return [
                Document(page_content="第一页", metadata={"source": self.path, "page": 0}),
                Document(page_content="第二页", metadata={"source": self.path, "page": 1}),
            ]

    monkeypatch.setattr(document_loader, "PyPDFLoader", FakeLoader)
    docs = document_loader.load_pdf(str(tmp_path / "a.pdf"), source_name="报告.pdf")

    assert docs[0].metadata["source"] == "报告.pdf"
    assert docs[0].metadata["page"] == 1
    assert docs[1].metadata["page"] == 2


def test_load_document_unsupported_type(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        document_loader.load_document(str(path))


def test_load_document_dispatches_by_suffix(tmp_path, monkeypatch):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    docx = tmp_path / "b.docx"
    docx.write_bytes(b"")

    monkeypatch.setattr(document_loader, "load_pdf", lambda p, source_name=None: ["pdf"])
    monkeypatch.setattr(document_loader, "load_docx", lambda p, source_name=None: ["docx"])

    assert document_loader.load_document(str(pdf)) == ["pdf"]
    assert document_loader.load_document(str(docx)) == ["docx"]
