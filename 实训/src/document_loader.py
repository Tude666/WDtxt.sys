"""文档解析：将 PDF / Word(.docx) 解析为 LangChain Document 列表。

解析后的每个 Document 统一携带规范 metadata：
  - source：来源文件名（友好名，可传入 source_name 覆盖默认的临时路径名）
  - page：页码（PDF 从 1 起；docx 无真实分页，不设 page）
"""
from pathlib import Path

import config  # noqa: F401  确保 HF_ENDPOINT 在 langchain_community 导入前生效

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document


def load_pdf(file_path: str | Path, source_name: str | None = None) -> list[Document]:
    """解析 PDF 文件，按页返回 Document，并补上 source / page 元数据。"""
    loader = PyPDFLoader(str(file_path))
    docs = loader.load()
    name = source_name or Path(file_path).name
    for i, doc in enumerate(docs, 1):
        doc.metadata["source"] = name
        doc.metadata["page"] = i  # 归一为 1 起的人类可读页码
    return docs


def load_docx(file_path: str | Path, source_name: str | None = None) -> list[Document]:
    """解析 Word(.docx) 文件，读取段落文本，返回单个 Document。

    注：仅支持 .docx（Office 2007+），不支持老式 .doc 格式；
    python-docx 不追踪真实分页，故不设 page 元数据。
    """
    doc = DocxDocument(str(file_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    name = source_name or Path(file_path).name
    return [Document(page_content=text, metadata={"source": name})]


def load_document(file_path: str | Path, source_name: str | None = None) -> list[Document]:
    """根据扩展名分发到对应解析器，source_name 用于标记友好来源文件名。"""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return load_pdf(file_path, source_name)
    if suffix == ".docx":
        return load_docx(file_path, source_name)
    raise ValueError(f"不支持的文件类型：{suffix}（仅支持 .pdf 和 .docx）")
